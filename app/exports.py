import re
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def _safe_filename(value: str) -> str:
    cleaned = re.sub(r"[^\w\s-]", "", value, flags=re.UNICODE).strip()
    cleaned = re.sub(r"[-\s]+", "-", cleaned)
    return cleaned[:80] or "event"


def build_event_registrations_workbook(event, registrations) -> BytesIO:
    from app.fields import visible_columns

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Registrations"

    columns = visible_columns(event.settings) + [("registered_by", "Registered by")]
    headers = [label for _key, label in columns] + ["Registered at (UTC)"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for reg in registrations:
        row = [reg.column_value(key) for key, _label in columns]
        row.append(reg.created_at.strftime("%Y-%m-%d %H:%M:%S"))
        sheet.append(row)

    for index in range(1, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(index)].width = 22

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def event_export_filename(event) -> str:
    date_part = event.event_date.strftime("%Y-%m-%d")
    return f"{_safe_filename(event.name)}-{date_part}-registrations.xlsx"
